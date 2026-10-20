"""Builds the Signal Hunter report (Word + JSON) from the saved state files only.
Numbers come from the provenance store and the calculation log; text sections come from the model's
narrative, which may only cite fact ids (F###) and calculation ids (C###)."""
import json

import scoring
from agent import DATA, OUT

ELEMENTS = [("discovery", "1. The Discovery"), ("why_it_matters", "2. Why It Matters"),
            ("evidence", "3. Evidence"), ("timeline", "4. Timeline"), ("numbers", "5. Numbers"),
            ("counterargument", "6. Counterargument"), ("falsification", "7. Falsification Condition"),
            ("evidence_strength", "8. Evidence Strength")]


def _load(name, default):
    p = DATA / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default


def state():
    s = dict(run=_load("run.json", {}), narrative=_load("narrative.json", {}),
             documents=_load("documents.json", []), facts=_load("facts_company.json", []),
             calcs=_load("calc_log.json", []), relationships=_load("relationships.json", {}),
             scores=_load("scores.json", {}), leads=_load("leads.json", []))
    n = s["narrative"]
    qc = {"eight_elements_complete": all(n.get(k) for k, _ in ELEMENTS),
          "every_number_has_source_tag": bool(s["facts"]) and all(f.get("doc") and f.get("page") for f in s["facts"]),
          "terms_defined_plain_language": bool(n.get("terms")),
          "calculation_log_present": bool(s["calcs"]),
          "verification_pass_done": bool(s["facts"]) and all("verified" in f for f in s["facts"]),
          "same_checks_run_on_peers": bool(n.get("peers_checked")),
          "applicability_and_limitations_stated": bool(n.get("limitations")),
          "coverage_grade_stated": bool(n.get("coverage_grade"))}
    s["run_qc"] = qc
    s["run_grade"] = scoring.run_grade(qc)
    return s


def build_json(path):
    s = state()
    OUT.mkdir(exist_ok=True)
    open(path, "w", encoding="utf-8").write(json.dumps(s, indent=2, ensure_ascii=False))
    return path


def _table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)
    return t


def _text(doc, value):
    if isinstance(value, list):
        for v in value:
            doc.add_paragraph(str(v), style="List Bullet")
    elif value:
        doc.add_paragraph(str(value))
    else:
        doc.add_paragraph("Not produced in this run.")


def build_docx(path):
    from docx import Document
    from docx.shared import Pt
    s = state()
    run, n = s["run"], s["narrative"]
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(f"Signal Hunter Report: {run.get('company', '?')}", 0)
    verdict = s["scores"].get("verdict", "leads not scored")
    _table(doc, ["Item", "Value"], [
        ["Company type (archetype)", run.get("archetype", "?")],
        ["As-of date / mode", f"{run.get('as_of', '?')} / {run.get('mode', '?')}"],
        ["Outcome (computed by scoring.py)", verdict],
        ["Run quality grade (code)", s["run_grade"]],
        ["Coverage grade", n.get("coverage_grade", "not stated")],
        ["Peers used", n.get("peers_checked") or "none (not checked in this run)"],
        ["Run cost (estimated)", f"US$ {run.get('cost', {}).get('estimated_usd', '?')}"],
    ])
    if run.get("stopped_early"):
        doc.add_paragraph(f"Note: the run stopped early ({run['stopped_early']}). Sections below show what was "
                          "completed; missing parts are marked.")

    doc.add_heading("Eight required elements", 1)
    for key, title in ELEMENTS:
        doc.add_heading(title, 2)
        if key == "numbers":
            used = n.get("numbers_calc_ids") or [c["id"] for c in s["calcs"]]
            rows = [[c["id"], c.get("label") or c["name"], c["value"], c["formula"], ", ".join(c["fact_ids"])]
                    for c in s["calcs"] if c["id"] in used]
            _table(doc, ["Calc", "Metric", "Value (code)", "Formula", "Fact ids"], rows) if rows else \
                doc.add_paragraph("No calculations were run.")
        else:
            _text(doc, n.get(key))

    doc.add_heading("Lead Register (scored by code)", 1)
    rows = s["scores"].get("rows", [])
    if rows:
        _table(doc, ["Lead", "Description", "Score", "Outcome", "Evidence label", "Gates"],
               [[r["id"], r["name"], r["final"], r["outcome"], r["evidence_label"],
                 "PASS" if r["gates_passed"] else "fail: " + ", ".join(x.split("_")[0] for x in r["failed"])]
                for r in rows])
    else:
        doc.add_paragraph("No leads were registered.")

    doc.add_heading("Relationship Map (R1 to R17, computed by code)", 1)
    edges = s["relationships"].get("edges", [])
    if edges:
        _table(doc, ["Edge", "Relationship", "Status", "Detail"],
               [[e["id"], e["relationship"], e["status"], e["detail"]] for e in edges])
    else:
        doc.add_paragraph("Relationships were not evaluated.")

    doc.add_heading("Evidence Trail (provenance store)", 1)
    _table(doc, ["Fact", "Metric", "Value", "Unit", "Period", "Scope", "Doc", "Page", "Verified by code"],
           [[f["id"], f.get("metric"), f.get("value"), f.get("unit"), f.get("period"), f.get("scope"),
             f.get("doc"), f.get("page"), "yes" if f.get("verified") else "NO: " + f.get("verify_note", "")]
            for f in s["facts"]])

    doc.add_heading("Source List and Coverage Log", 1)
    _table(doc, ["Doc", "Need", "Quarter", "Tier", "Family", "Found by", "URL"],
           [[d["id"], d.get("need"), d.get("quarter"), d.get("tier"), d.get("family"), d.get("route"), d.get("url")]
            for d in s["documents"]])
    for k in ("missing_sources", "limitations"):
        if n.get(k):
            doc.add_heading(k.replace("_", " ").capitalize(), 2)
            _text(doc, n[k])

    doc.add_heading("Verification Log", 1)
    nv = [f for f in s["facts"] if not f.get("verified")]
    doc.add_paragraph(f"{len(s['facts']) - len(nv)} of {len(s['facts'])} facts were re-found on the cited page "
                      "by code (excerpt words and the number). Unverified facts are excluded from relationships "
                      "and make gate F fail for any lead that uses them.")
    if n.get("verification_notes"):
        _text(doc, n["verification_notes"])

    if n.get("terms"):
        doc.add_heading("Terms used", 1)
        _text(doc, n["terms"])
    doc.add_heading("Run QC (code)", 1)
    _table(doc, ["Check", "Passed"], [[k, "yes" if v else "no"] for k, v in s["run_qc"].items()])
    doc.add_paragraph("This report is research, not investment advice. It makes no forecast and no allegation.")
    OUT.mkdir(exist_ok=True)
    doc.save(path)
    return path
